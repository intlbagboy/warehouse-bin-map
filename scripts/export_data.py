"""
Regenerate data/locations.json from the Warehouse Relocation Tracker workbook.

Usage:
    python3 scripts/export_data.py path/to/Warehouse_Relocation_Tracker.xlsx

Run this any time you update the Excel file, then commit + push data/locations.json
to redeploy the site with fresh data.
"""
import sys
import json
import datetime
from pathlib import Path

import openpyxl

SPOTS_PER_AISLE = 22
LEVELS_PER_AISLE = 4


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    aisle_codes = []
    for row in wb["Aisle List"].iter_rows(min_row=2, max_row=21, values_only=True):
        if row[0]:
            aisle_codes.append(str(row[0]))

    blocked_spots = set()
    if "Spots Blocked by Beams" in wb.sheetnames:
        for row in wb["Spots Blocked by Beams"].iter_rows(min_row=2, values_only=True):
            if row[0]:
                blocked_spots.add(str(row[0]).strip())

    # location code -> list of transfer log entries
    entries_by_loc = {}
    for row in wb["Transfer Log"].iter_rows(min_row=2, values_only=True):
        lp, sku, desc, qty, uom, loc, by, verified = (
            row[1], row[2], row[3], row[4], row[5], row[6], row[7], row[8]
        )
        if not lp or not loc:
            continue
        entries_by_loc.setdefault(loc, []).append({
            "licensePlate": lp,
            "sku": sku,
            "description": desc,
            "qty": qty,
            "uom": uom,
            "transferredBy": by,
            "verified": (verified == "Y"),
        })

    locations = {}
    for aisle in aisle_codes:
        for spot in range(1, SPOTS_PER_AISLE + 1):
            spot_key = f"{aisle}-{spot}"
            is_blocked = spot_key in blocked_spots
            for level in range(1, LEVELS_PER_AISLE + 1):
                code = f"{aisle}-{str(spot).zfill(2)}-{level}"
                if is_blocked:
                    locations[code] = {"status": "blocked", "entries": [], "qty": 0}
                    continue
                entries = entries_by_loc.get(code, [])
                if not entries:
                    locations[code] = {"status": "empty", "entries": [], "qty": 0}
                    continue
                needs_review = any(not e["verified"] for e in entries)
                total_qty = sum(e["qty"] or 0 for e in entries)
                locations[code] = {
                    "status": "review" if needs_review else "occupied",
                    "entries": entries,
                    "qty": total_qty,
                }

    return {
        "generatedAt": datetime.datetime.now().isoformat(timespec="seconds"),
        "aisles": [
            {"code": a, "spots": SPOTS_PER_AISLE, "levels": LEVELS_PER_AISLE}
            for a in aisle_codes
        ],
        "locations": locations,
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python3 scripts/export_data.py <path-to-workbook.xlsx>")
        sys.exit(1)

    data = load(sys.argv[1])
    out_path = Path(__file__).resolve().parent.parent / "data" / "locations.json"
    out_path.parent.mkdir(exist_ok=True)
    out_path.write_text(json.dumps(data, indent=2))
    print(f"Wrote {out_path} — {len(data['locations'])} locations across {len(data['aisles'])} aisle(s)")
